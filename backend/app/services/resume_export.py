from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


def _clean_lines(content: str) -> list[str]:
    return [
        line.strip()
        for line in content.splitlines()
    ]


def _is_section_heading(line: str) -> bool:
    """
    Simple first-version heading detection.

    Examples:
    EDUCATION
    EXPERIENCE
    PROJECTS
    TECHNICAL SKILLS
    """

    if not line:
        return False

    if len(line) > 60:
        return False

    letters = [
        char
        for char in line
        if char.isalpha()
    ]

    if not letters:
        return False

    return line.upper() == line


def build_resume_docx(content: str) -> bytes:
    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    lines = _clean_lines(content)

    first_text_line = True

    for line in lines:

        if not line:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(line)

        if first_text_line:
            run.bold = True
            run.font.size = Pt(16)

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            paragraph.paragraph_format.space_after = Pt(4)

            first_text_line = False

        elif _is_section_heading(line):
            run.bold = True
            run.font.size = Pt(10.5)

            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)

        else:
            run.font.size = Pt(10)

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


def build_resume_pdf(content: str) -> bytes:
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title="Resume",
    )

    title_style = ParagraphStyle(
        name="ResumeTitle",
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=18,
        alignment=TA_CENTER,
        spaceAfter=6,
    )

    heading_style = ParagraphStyle(
        name="ResumeHeading",
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=12,
        spaceBefore=6,
        spaceAfter=2,
    )

    body_style = ParagraphStyle(
        name="ResumeBody",
        fontName="Helvetica",
        fontSize=10,
        leading=12,
        spaceAfter=2,
    )

    story = []

    first_text_line = True

    for line in _clean_lines(content):

        if not line:
            story.append(Spacer(1, 4))
            continue

        # Paragraph supports markup, so escape user text first.
        safe_line = escape(line)

        if first_text_line:
            story.append(
                Paragraph(
                    safe_line,
                    title_style,
                )
            )

            first_text_line = False

        elif _is_section_heading(line):
            story.append(
                Paragraph(
                    safe_line,
                    heading_style,
                )
            )

        else:
            story.append(
                Paragraph(
                    safe_line,
                    body_style,
                )
            )

    document.build(story)

    output.seek(0)

    return output.getvalue()