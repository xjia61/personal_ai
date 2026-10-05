from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_resume_text(
    content: bytes,
    filename: str,
) -> str:
    filename = filename.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(BytesIO(content))

        if reader.is_encrypted:
            raise ValueError("Encrypted PDF is not supported")

        if len(reader.pages) > 20:
            raise ValueError("PDF exceeds the 20-page limit")

        parts = [
            page.extract_text() or ""
            for page in reader.pages
        ]

    elif filename.endswith(".docx"):
        document = Document(BytesIO(content))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text for cell in row.cells
                    )
                )

    else:
        raise ValueError(
            "Only PDF and DOCX files are supported"
        )

    text = "\n".join(parts).strip()

    if not text:
        raise ValueError(
            "No readable text was found in this file"
        )

    return text