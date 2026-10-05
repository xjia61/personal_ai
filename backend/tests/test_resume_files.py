from io import BytesIO

import pytest
from docx import Document

from app.services.resume_files import extract_resume_text


def test_docx_import():
    document = Document()

    document.add_paragraph("Software Engineer")
    document.add_paragraph("Python, React, SQL")

    buffer = BytesIO()
    document.save(buffer)

    text = extract_resume_text(
        buffer.getvalue(),
        "resume.docx",
    )

    assert "Software Engineer" in text
    assert "Python, React, SQL" in text


def test_unsupported_file():
    with pytest.raises(ValueError):
        extract_resume_text(
            b"hello",
            "resume.exe",
        )